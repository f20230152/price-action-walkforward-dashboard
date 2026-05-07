from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parent
OUT_DIR = BASE_DIR / "outputs" / "clean_walkforward"
XLSX_PATH = OUT_DIR / "walkforward_report.xlsx"


def read_csv(name: str, **kwargs) -> pd.DataFrame:
    path = OUT_DIR / name
    return pd.read_csv(path, **kwargs) if path.exists() else pd.DataFrame()


def autosize(ws, max_width: int = 60) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), max_width)


def style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autosize(ws)


def main() -> None:
    metrics = read_csv("walkforward_metrics.csv")
    decisions = read_csv("walkforward_decisions_all.csv")
    trades = read_csv("walkforward_trades_all.csv")
    daily = read_csv("walkforward_daily_pnl_all.csv")
    universe = read_csv("parameter_universe.csv")
    config = read_csv("research_config.csv")
    rankings = read_csv("walkforward_train_rankings_all.csv")

    best = metrics.sort_values(["daily_sharpe", "total_pnl_cents"], ascending=False).iloc[0]
    summary = pd.DataFrame(
        [
            ["Recommended selector", best["selector_objective"]],
            ["WF Net PnL (cents)", best["total_pnl_cents"]],
            ["WF Sharpe", best["daily_sharpe"]],
            ["WF Max Drawdown (cents)", best["max_drawdown_cents"]],
            ["WF Trades", best["trades"]],
            ["WF Trades/Day", best["trades_per_day"]],
            ["WF Win Rate", best["win_rate"]],
            ["WF Profit Factor", best["profit_factor"]],
            ["Candidate universe size", int(config["candidate_count"].iloc[0]) if not config.empty else ""],
            ["Train/Test", "3 months train / 1 month test"],
            ["Cost", "2c per side / 4c round trip"],
            ["Trade cap", "1 trade/day"],
        ],
        columns=["Metric", "Value"],
    )

    equity = daily.copy()
    if not equity.empty:
        date_col = equity.columns[0]
        equity[date_col] = pd.to_datetime(equity[date_col])
        for col in equity.columns[1:]:
            equity[f"{col}_equity"] = equity[col].fillna(0).cumsum()

    with pd.ExcelWriter(XLSX_PATH, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        metrics.to_excel(writer, sheet_name="WF Metrics", index=False)
        decisions.to_excel(writer, sheet_name="WF Decisions", index=False)
        trades.to_excel(writer, sheet_name="WF Trades", index=False)
        daily.to_excel(writer, sheet_name="Daily PnL", index=False)
        equity.to_excel(writer, sheet_name="Equity Curves", index=False)
        config.to_excel(writer, sheet_name="Research Config", index=False)
        universe.to_excel(writer, sheet_name="Parameter Universe", index=False)
        rankings.to_excel(writer, sheet_name="Train Rankings", index=False)

    wb = load_workbook(XLSX_PATH)
    for ws in wb.worksheets:
        style_sheet(ws)

    if "Equity Curves" in wb.sheetnames:
        ws = wb["Equity Curves"]
        if ws.max_row > 2 and ws.max_column > 4:
            chart = LineChart()
            chart.title = "Walk-Forward Equity Curves"
            chart.y_axis.title = "Cumulative cents"
            chart.x_axis.title = "Date"
            data = Reference(ws, min_col=5, max_col=ws.max_column, min_row=1, max_row=ws.max_row)
            cats = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            chart.height = 12
            chart.width = 24
            ws.add_chart(chart, "J2")

    if "WF Decisions" in wb.sheetnames:
        ws = wb["WF Decisions"]
        if ws.max_row > 2:
            chart = BarChart()
            chart.title = "Test Month PnL By Selector"
            chart.y_axis.title = "Cents"
            data = Reference(ws, min_col=11, min_row=1, max_row=ws.max_row)
            cats = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            chart.height = 10
            chart.width = 20
            ws.add_chart(chart, "P2")

    wb.save(XLSX_PATH)
    # Verification load.
    check = load_workbook(XLSX_PATH, data_only=False, read_only=True)
    required = {"Summary", "WF Metrics", "WF Decisions", "WF Trades", "Daily PnL"}
    missing = required - set(check.sheetnames)
    if missing:
        raise RuntimeError(f"Missing sheets: {sorted(missing)}")
    print(XLSX_PATH)


if __name__ == "__main__":
    main()
